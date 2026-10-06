from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path('/home/ubuntu/keywordslot')
CATALOG = ROOT / 'src/data/keywordCatalog.json'
SALES_DETAIL = Path('/home/ubuntu/output/판매키워드_날짜별_정제.xlsx')
SPLIT_AUDIT = Path('/home/ubuntu/output/결합키워드_분리내역.json')
REPORT = Path('/home/ubuntu/output/판매키워드_판매날짜_연결검증.json')


def normalize_keyword(value: Any) -> str:
    return re.sub(r'\s+', ' ', str(value or '')).strip(' ,.')


def iso_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raise ValueError(f'Unsupported date value: {value!r}')


def main() -> None:
    catalog: list[dict[str, Any]] = json.loads(CATALOG.read_text(encoding='utf-8'))
    audit = json.loads(SPLIT_AUDIT.read_text(encoding='utf-8'))
    split_map = {
        normalize_keyword(entry['sourceKeyword']): [normalize_keyword(part['keyword']) for part in entry['splitKeywords']]
        for entry in audit['entries']
    }

    workbook = load_workbook(SALES_DETAIL, data_only=True, read_only=True)
    worksheet = workbook['키워드 상세']
    dates_by_keyword: dict[str, list[str]] = defaultdict(list)
    for sale_date, _timestamp, keyword, _source_row in worksheet.iter_rows(min_row=2, values_only=True):
        normalized = normalize_keyword(keyword)
        if not normalized or not sale_date:
            continue
        sale_day = iso_date(sale_date)
        for part in split_map.get(normalized, [normalized]):
            dates_by_keyword[part].append(sale_day)

    final_catalog: list[dict[str, Any]] = []
    missing: list[str] = []
    count_mismatches: list[dict[str, Any]] = []
    for item in catalog:
        keyword = normalize_keyword(item['keyword'])
        occurrences = sorted(dates_by_keyword.get(keyword, []))
        if not occurrences:
            missing.append(keyword)
            continue
        if len(occurrences) != int(item['salesCount']):
            count_mismatches.append({
                'keyword': keyword,
                'catalogSalesCount': int(item['salesCount']),
                'datedOccurrences': len(occurrences),
            })
            continue
        sale_dates = sorted(set(occurrences))
        sale_months = sorted({f'{int(day[5:7])}월' for day in sale_dates}, key=lambda value: int(value[:-1]))
        final_catalog.append({
            'keyword': keyword,
            'englishKeyword': item['englishKeyword'],
            'salesCount': int(item['salesCount']),
            'salesDates': sale_dates,
            'salesMonths': sale_months,
        })

    report = {
        'catalogKeywords': len(catalog),
        'finalKeywords': len(final_catalog),
        'catalogSalesMentions': sum(int(item['salesCount']) for item in catalog),
        'datedSalesMentions': sum(int(item['salesCount']) for item in final_catalog),
        'missingKeywords': missing,
        'countMismatches': count_mismatches,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    if missing or count_mismatches:
        raise RuntimeError(f'Sales-date mapping incomplete: missing={len(missing)}, mismatches={len(count_mismatches)}')

    CATALOG.write_text(json.dumps(final_catalog, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({
        'keywords': len(final_catalog),
        'salesMentions': sum(item['salesCount'] for item in final_catalog),
        'dateRange': [min(day for item in final_catalog for day in item['salesDates']), max(day for item in final_catalog for day in item['salesDates'])],
        'months': sorted({month for item in final_catalog for month in item['salesMonths']}, key=lambda value: int(value[:-1])),
    }, ensure_ascii=False))


if __name__ == '__main__':
    main()
