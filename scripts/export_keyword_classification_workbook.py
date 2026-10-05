from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path('/home/ubuntu/keywordslot')
CATALOG = ROOT / 'src/data/keywordCatalog.json'
AUDIT = Path('/home/ubuntu/output/결합키워드_분리내역.json')
OUTPUT = Path('/home/ubuntu/output/판매키워드_슬롯카테고리_분류.xlsx')
ORDER = ['자연·식물', '동물·곤충', '사람·생활', '음식·음료', '기념일·시즌', '교육·사무', '교통·여행', '건강·안전', '문화·취미', '비즈니스·금융', '종교·전통', '장식·추상']
DESCRIPTION = {
    '자연·식물': '꽃, 나무, 날씨, 지형, 과일·채소 등',
    '동물·곤충': '동물, 새, 물고기, 곤충 등',
    '사람·생활': '인물, 가족, 일상, 생활 물품 등',
    '음식·음료': '음식, 식재료, 음료, 조리 도구 등',
    '기념일·시즌': '명절, 계절 행사, 기념일, 축제 등',
    '교육·사무': '학교, 독서, 문서, 학습·사무 도구 등',
    '교통·여행': '이동수단, 도로, 지역, 관광 등',
    '건강·안전': '의료, 운동, 위생, 안전, 재난 등',
    '문화·취미': '음악, 스포츠, 예술, 놀이, 미디어 등',
    '비즈니스·금융': '경제, 쇼핑, 업무, 금융, 디지털 서비스 등',
    '종교·전통': '종교, 전통 문화, 국가 상징, 역사 등',
    '장식·추상': '아이콘, 패턴, 감정, 기호, 추상 요소 등',
}


def style_sheet(sheet, widths: dict[str, int]) -> None:
    navy, white = '1F4E78', 'FFFFFF'
    line = Side(style='thin', color='D9E2F3')
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 26
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=line)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=f'A1:{chr(64 + sheet.max_column)}{sheet.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def main() -> None:
    catalog = json.loads(CATALOG.read_text(encoding='utf-8'))
    audit = json.loads(AUDIT.read_text(encoding='utf-8'))
    count_by_category = Counter(item['category'] for item in catalog)
    sales_by_category = Counter()
    for item in catalog:
        sales_by_category[item['category']] += item['salesCount']

    workbook = Workbook()
    summary = workbook.active
    summary.title = '카테고리 요약'
    summary.append(['카테고리', '고유 키워드 수', '판매 언급 합계', '설명'])
    for category in ORDER:
        summary.append([category, count_by_category[category], sales_by_category[category], DESCRIPTION[category]])
    summary.append(['합계', len(catalog), sum(item['salesCount'] for item in catalog), '결합 키워드 11개를 원자 키워드로 분리·재분류한 최신 카탈로그 기준'])
    style_sheet(summary, {'A': 20, 'B': 16, 'C': 16, 'D': 62})
    add_table(summary, 'CategorySummaryTable')

    details = workbook.create_sheet('키워드 상세')
    details.append(['키워드', '영문 프롬프트 키워드', '카테고리', '난이도', '판매 언급 수'])
    for item in sorted(catalog, key=lambda item: (ORDER.index(item['category']), -item['salesCount'], item['keyword'])):
        details.append([item['keyword'], item['englishKeyword'], item['category'], item['difficulty'], item['salesCount']])
    style_sheet(details, {'A': 30, 'B': 34, 'C': 20, 'D': 12, 'E': 16})
    add_table(details, 'KeywordCategoryTable')

    splits = workbook.create_sheet('결합 키워드 분리 내역')
    splits.append(['원본 결합 키워드', '원본 판매 언급', '분리 키워드', '카테고리', '난이도', '영문 프롬프트 키워드'])
    for entry in audit['entries']:
        for split in entry['splitKeywords']:
            splits.append([
                entry['sourceKeyword'], entry['sourceSalesCount'], split['keyword'], split['category'],
                split['difficulty'], split['englishKeyword'],
            ])
    style_sheet(splits, {'A': 58, 'B': 16, 'C': 22, 'D': 20, 'E': 12, 'F': 34})
    add_table(splits, 'CompositeKeywordSplitTable')
    for row in range(2, splits.max_row + 1):
        splits.row_dimensions[row].height = 32

    workbook.properties.title = '판매 키워드 슬롯 카테고리 분류'
    workbook.properties.subject = '결합 키워드 분리 및 카테고리 재분류'
    workbook.save(OUTPUT)
    print(OUTPUT)
    print('sheets=', workbook.sheetnames, 'keywords=', len(catalog), 'splits=', len(audit['entries']))


if __name__ == '__main__':
    main()
