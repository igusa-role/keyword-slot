from __future__ import annotations

import concurrent.futures as cf
import json
import time
from collections import Counter
from pathlib import Path
from typing import Any

from openai import OpenAI
from openpyxl import load_workbook

SOURCE = Path('/home/ubuntu/output/판매키워드_날짜별_정제.xlsx')
DESTINATION = Path('/home/ubuntu/keywordslot/src/data/keywordCatalog.json')
CACHE = Path('/home/ubuntu/keywordslot/.keyword_catalog_cache.json')
MODEL = 'gpt-5-mini'
BATCH_SIZE = 100
MAX_WORKERS = 5

CATEGORIES = [
    '자연·식물', '동물·곤충', '사람·생활', '음식·음료', '기념일·시즌', '교육·사무',
    '교통·여행', '건강·안전', '문화·취미', '비즈니스·금융', '종교·전통', '장식·추상',
]
DIFFICULTIES = ['쉬움', '보통', '어려움']

SYSTEM_PROMPT = '''당신은 한국어 그림 소재 키워드 분류 전문가입니다. 입력된 판매 키워드를 아래의 정확히 12개 카테고리와 3개 난이도로 분류하세요.

카테고리(반드시 하나):
- 자연·식물: 꽃, 나무, 날씨, 계절 자연물, 지형, 과일·채소 등
- 동물·곤충: 동물, 새, 물고기, 곤충 등
- 사람·생활: 인물, 가족, 일상 행동, 의복, 주거, 생활 물품 등
- 음식·음료: 음식, 식재료, 음료, 조리·식기 등
- 기념일·시즌: 명절, 계절·월별 행사, 기념일, 축제, 졸업 등
- 교육·사무: 학교, 독서, 문서, 사무, 학습 도구 등
- 교통·여행: 이동수단, 도로, 지역, 관광, 여행 등
- 건강·안전: 의료, 운동, 위생, 안전, 재난 등
- 문화·취미: 음악, 스포츠, 예술, 놀이, 미디어, 취미 등
- 비즈니스·금융: 경제, 쇼핑, 업무, 금융, 디지털 서비스 등
- 종교·전통: 종교, 전통 문화, 국가 상징, 역사 등
- 장식·추상: 하트, 아이콘, 패턴, 감정, 기호, 추상 요소 등

난이도(반드시 하나):
- 쉬움: 단일하고 친숙한 소재, 간단한 아이콘/윤곽선으로 그릴 수 있음
- 보통: 간단한 복합 소재, 행사명, 동작이 있는 대상
- 어려움: 여러 대상 관계, 전문 소재, 복잡한 장면 또는 추상적 개념

키워드의 판매 빈도와 원문을 변경하지 마세요. 반드시 모든 id를 정확히 한 번씩 반환하세요.'''

SCHEMA: dict[str, Any] = {
    'type': 'json_schema',
    'json_schema': {
        'name': 'keyword_categories',
        'strict': True,
        'schema': {
            'type': 'object',
            'properties': {
                'items': {
                    'type': 'array',
                    'items': {
                        'type': 'object',
                        'properties': {
                            'id': {'type': 'integer'},
                            'category': {'type': 'string', 'enum': CATEGORIES},
                            'difficulty': {'type': 'string', 'enum': DIFFICULTIES},
                        },
                        'required': ['id', 'category', 'difficulty'],
                        'additionalProperties': False,
                    },
                },
            },
            'required': ['items'],
            'additionalProperties': False,
        },
    },
}


def load_keywords() -> list[dict[str, Any]]:
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    worksheet = workbook['키워드 상세']
    counts = Counter(str(row[2]).strip() for row in worksheet.iter_rows(min_row=2, values_only=True) if row[2])
    return [
        {'id': index, 'keyword': keyword, 'salesCount': count}
        for index, (keyword, count) in enumerate(sorted(counts.items()), start=1)
    ]


def classify_batch(batch: list[dict[str, Any]]) -> dict[int, dict[str, str]]:
    expected_ids = {item['id'] for item in batch}
    prompt_data = [{'id': item['id'], 'keyword': item['keyword']} for item in batch]
    last_error: Exception | None = None
    for attempt in range(4):
        try:
            response = OpenAI().chat.completions.create(
                model=MODEL,
                messages=[
                    {'role': 'system', 'content': SYSTEM_PROMPT},
                    {'role': 'user', 'content': '다음 JSON 배열의 모든 키워드를 분류하세요.\n' + json.dumps(prompt_data, ensure_ascii=False)},
                ],
                response_format=SCHEMA,
                max_completion_tokens=6500,
                extra_body={'reasoning': {'effort': 'minimal'}},
            )
            data = json.loads(response.choices[0].message.content)
            ids = [item['id'] for item in data['items']]
            if set(ids) != expected_ids or len(ids) != len(expected_ids):
                raise ValueError('Classification response IDs did not match batch IDs')
            return {item['id']: {'category': item['category'], 'difficulty': item['difficulty']} for item in data['items']}
        except Exception as error:
            last_error = error
            time.sleep(2 ** attempt)
    raise RuntimeError(f'Classification failed after retries: {last_error}')


def main() -> None:
    keywords = load_keywords()
    if CACHE.exists():
        cached = json.loads(CACHE.read_text(encoding='utf-8'))
        if cached.get('keywords') == keywords:
            classifications = {int(key): value for key, value in cached['classifications'].items()}
            print('Using cached classifications')
        else:
            classifications = {}
    else:
        classifications = {}

    missing = [item for item in keywords if item['id'] not in classifications]
    if missing:
        batches = [missing[index:index + BATCH_SIZE] for index in range(0, len(missing), BATCH_SIZE)]
        with cf.ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = [executor.submit(classify_batch, batch) for batch in batches]
            for index, future in enumerate(cf.as_completed(futures), start=1):
                classifications.update(future.result())
                print(f'Classified {index}/{len(batches)} batches')
        CACHE.write_text(json.dumps({'keywords': keywords, 'classifications': classifications}, ensure_ascii=False), encoding='utf-8')

    catalog = [
        {
            'keyword': item['keyword'],
            'salesCount': item['salesCount'],
            **classifications[item['id']],
        }
        for item in keywords
    ]
    DESTINATION.parent.mkdir(parents=True, exist_ok=True)
    DESTINATION.write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding='utf-8')
    counts = Counter(item['category'] for item in catalog)
    print(f'Wrote {DESTINATION} with {len(catalog)} unique keywords')
    print(json.dumps(dict(sorted(counts.items())), ensure_ascii=False))


if __name__ == '__main__':
    main()
